"""安装包自检。

``EngineerDiff.exe --self-test`` 在无界面情况下验证：核心比较引擎、Excel/CSV
读写、内置 OCR 模型、图片比较流程。不依赖 Office、网络、外部 Python。
"""

from __future__ import annotations

import json
import os
import sys
import tempfile
import time
import traceback
from pathlib import Path
from typing import Any, Callable

from . import version

REPORT_NAME = "self-test-report.json"


def _safe_write(path: Path, payload: dict[str, Any]) -> None:
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    except OSError:
        pass


def _make_ocr_images(folder: Path) -> tuple[Path, Path]:
    from PIL import Image, ImageDraw, ImageFont

    def font(size: int):
        try:
            return ImageFont.load_default(size=size)
        except TypeError:  # Pillow < 10.1
            return ImageFont.load_default()

    img_a = Image.new("RGB", (1000, 360), "white")
    da = ImageDraw.Draw(img_a)
    da.text((40, 50), "VERSION 1.2.3", fill="black", font=font(36))
    da.text((40, 130), "CURRENT 1.50 A", fill="black", font=font(36))
    da.text((40, 210), "STATUS ON", fill="black", font=font(36))
    path_a = folder / "device-a.png"
    img_a.save(path_a)

    img_b = Image.new("RGB", (800, 288), "white")
    db = ImageDraw.Draw(img_b)
    db.text((32, 40), "VERSION 1.2.4", fill="black", font=font(29))
    db.text((32, 104), "CURRENT 1.80 A", fill="black", font=font(29))
    db.text((32, 168), "STATUS ON", fill="black", font=font(29))
    path_b = folder / "device-b.png"
    img_b.save(path_b)
    return path_a, path_b


def run_self_test(report_path: str | Path = REPORT_NAME) -> int:
    started = time.time()
    checks: list[dict[str, Any]] = []

    def check(name: str, fn: Callable[[], str]) -> None:
        t0 = time.time()
        try:
            detail = fn()
            checks.append({"name": name, "ok": True, "detail": detail, "seconds": round(time.time() - t0, 2)})
        except Exception as exc:  # noqa: BLE001
            checks.append(
                {
                    "name": name,
                    "ok": False,
                    "detail": f"{type(exc).__name__}: {exc}",
                    "traceback": traceback.format_exc(limit=8),
                    "seconds": round(time.time() - t0, 2),
                }
            )

    with tempfile.TemporaryDirectory(prefix="engineerdiff-selftest-") as tmp:
        tmp_path = Path(tmp)

        def check_text() -> str:
            from .textdiff.engine import compare_texts

            cmp = compare_texts("软件版本 V1.2.3\n", "软件版本 V1.2.4\n")
            assert len(cmp.diffs) == 1, f"期望 1 处差异，实际 {len(cmp.diffs)}"
            spans = cmp.left.lines[0].spans
            assert len(spans) == 1 and (spans[0].start, spans[0].end) == (10, 11), "字符级差异定位错误"
            return "文本字符级差异、中文与差异定位正常"

        def check_csv() -> str:
            from .spreadsheet.csv_parser import parse_csv
            from .spreadsheet.engine import compare_workbooks
            from .models import Workbook

            left = tmp_path / "a.csv"
            right = tmp_path / "b.csv"
            left.write_text('设备,数值,备注\n"装置,一",1.50,"含,逗号"\n装置二,2.00,"多行\n字段"\n', encoding="utf-8")
            right.write_text('设备,数值,备注\n"装置,一",1.80,"含,逗号"\n装置二,2.00,"多行\n字段"\n', encoding="utf-8")
            ls, delim, _ = parse_csv(left)
            rs, _, _ = parse_csv(right)
            assert delim == ","
            assert len(ls.rows) >= 3
            cmp = compare_workbooks(
                Workbook(str(left), [ls], "csv"), Workbook(str(right), [rs], "csv")
            )
            assert len(cmp.diffs) == 1, f"CSV 期望 1 处差异，实际 {len(cmp.diffs)}"
            loc = cmp.diffs[0].left
            assert loc is not None and loc.col == 1
            return "CSV 引号/逗号/多行字段解析与单元格差异正常"

        def check_excel() -> str:
            import openpyxl
            from .spreadsheet.engine import compare_workbooks
            from .spreadsheet.loader import load_workbook

            def make(path: Path, version_text: str, reverse: bool) -> None:
                wb = openpyxl.Workbook()
                ws1 = wb.active
                ws1.title = "软件版本"
                ws1["A1"] = "装置"
                ws1["B1"] = "软件版本"
                ws1["A2"] = "装置一"
                ws1["B2"] = version_text
                ws2 = wb.create_sheet("保护定值")
                ws2["A1"] = "定值项"
                ws2["B1"] = "值"
                ws2["A2"] = "过流 I 段"
                ws2["B2"] = 1.50
                ws2["B3"] = "=B2*2"
                if reverse:
                    wb.move_sheet("保护定值", offset=-1)
                wb.save(path)

            a = tmp_path / "wb-a.xlsx"
            b = tmp_path / "wb-b.xlsx"
            make(a, "V1.2.3", False)
            make(b, "V1.2.4", True)
            cmp = compare_workbooks(load_workbook(a), load_workbook(b))
            by_name = {s.name: s for s in cmp.sheets}
            assert set(by_name) == {"软件版本", "保护定值"}, "多 Sheet 读取不完整"
            assert by_name["软件版本"].status.value == "different"
            assert by_name["保护定值"].status.value == "equal", "Sheet 顺序不同不应误报"
            assert len(cmp.diffs) == 1, f"Excel 期望 1 处差异，实际 {len(cmp.diffs)}"
            return "多 Sheet 名称配对、公式读取、单元格差异与顺序无关性正常"

        def check_ocr_models() -> str:
            from .ocr.engine import OcrService

            info = OcrService(cache_dir=tmp_path / "ocr-cache").model_info()
            assert info["models_present"], "安装包内未找到 OCR 模型"
            assert len(info["models"]) >= 3, f"OCR 模型不完整：{info['models']}"
            return f"{len(info['models'])} 个模型，{info['backend']}，离线 CPU"

        def check_image_pipeline() -> str:
            from .compare import run_comparison
            from .models import AppSettings, CompareMode
            from .ocr.engine import OcrService

            a, b = _make_ocr_images(tmp_path)
            settings = AppSettings()
            settings.image.confidence_threshold = 0.4
            svc = OcrService(cache_dir=tmp_path / "ocr-cache")
            result = run_comparison(CompareMode.IMAGE, a, b, settings, ocr_service=svc)
            img = result.image
            assert img is not None
            assert len(img.left.boxes) >= 2 and len(img.right.boxes) >= 2, (
                f"OCR 文字块不足：左 {len(img.left.boxes)} / 右 {len(img.right.boxes)}"
            )
            assert len(img.diffs) >= 1, "不同版本号的截图未检出差异"
            for d in img.diffs:
                if d.left is not None:
                    assert d.left.bbox is not None and d.left.bbox.w > 0, "左侧差异缺少坐标"
                if d.right is not None:
                    assert d.right.bbox is not None and d.right.bbox.w > 0, "右侧差异缺少坐标"
            for box in img.left.boxes:
                assert 0 <= box.bbox.x < img.left.width and 0 <= box.bbox.y < img.left.height
            for box in img.right.boxes:
                assert 0 <= box.bbox.x < img.right.width and 0 <= box.bbox.y < img.right.height
            return f"OCR 识别 {len(img.left.boxes)}+{len(img.right.boxes)} 块，检出 {len(img.diffs)} 处图片差异"

        def check_gui_imports() -> str:
            os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
            from PySide6.QtWidgets import QApplication

            from .ui.main_window import MainWindow
            from .ui.batch_dialog import BatchCompareDialog  # noqa: F401
            from .ui.image_view import ImagePane  # noqa: F401
            from .ui.spreadsheet_side import SpreadsheetSide  # noqa: F401

            app = QApplication.instance() or QApplication([])
            window = MainWindow()
            window.resize(1100, 720)
            window.show()
            app.processEvents()
            pixmap = window.grab()
            assert not pixmap.isNull() and pixmap.width() >= 900, "主窗口无法渲染"
            shot = tmp_path / "window-smoke.png"
            pixmap.save(str(shot), "PNG")
            window.close()
            app.processEvents()
            return f"PySide6/Qt 插件加载、主窗口渲染正常（截图 {shot.name} {pixmap.width()}×{pixmap.height()}）"

        def check_file_loaders() -> str:
            import xlrd  # noqa: F401
            import openpyxl  # noqa: F401
            from .file_loader import determine_mode
            from .models import CompareMode

            assert determine_mode(tmp_path / "a.xlsx", tmp_path / "b.xls") == CompareMode.EXCEL
            assert determine_mode(tmp_path / "a.png", tmp_path / "b.jpg") == CompareMode.IMAGE
            return "文件类型识别与 Office 读取组件可用"

        check("文本比较", check_text)
        check("CSV 比较", check_csv)
        check("Excel 比较", check_excel)
        check("OCR 模型内置", check_ocr_models)
        check("图片 OCR 比较", check_image_pipeline)
        check("界面组件导入", check_gui_imports)
        check("文件读取组件", check_file_loaders)

    ok = all(c["ok"] for c in checks)
    report = {
        "ok": ok,
        "version": version.__version__,
        "app": version.APP_NAME,
        "timestamp": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "python": sys.version.split()[0],
        "frozen": bool(getattr(sys, "frozen", False)),
        "elapsed_seconds": round(time.time() - started, 2),
        "checks": checks,
        "error": "" if ok else "自检存在失败项",
    }
    out = Path(report_path)
    if not out.is_absolute():
        out = Path.cwd() / out
    _safe_write(out, report)
    try:
        if sys.stdout is not None:
            print(json.dumps({"ok": ok, "report": str(out), "checks": checks}, ensure_ascii=False, indent=2))
    except Exception:
        pass
    return 0 if ok else 1
