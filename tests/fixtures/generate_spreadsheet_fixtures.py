"""生成 EngineerDiff 表格测试样本（XLSX / XLS / CSV）。

运行方式::

    .venv/bin/python tests/fixtures/generate_spreadsheet_fixtures.py

生成的样本会写入本目录（小文件，可以随测试一起提交）；测试开头也会调用
:func:`ensure_fixtures`，缺失时自动生成，因此删除样本后测试仍可重复运行。

XLSX 的公式缓存值：openpyxl 保存公式时不会写入计算结果。为了模拟“Excel 保存
过计算结果”的真实文件，本脚本在保存 zip 之后直接向 worksheet XML 的 ``<v>``
节点写入确定值（不依赖 Excel / LibreOffice）。
"""

from __future__ import annotations

import datetime as _dt
import io
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Any

FIXTURES_DIR = Path(__file__).resolve().parent

_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

ET.register_namespace("", _MAIN_NS)


# ---------------------------------------------------------------------------
# XLSX 公共
# ---------------------------------------------------------------------------


def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="999999")
    return {
        "title": (
            Font(name="Microsoft YaHei", size=16, bold=True, color="C00000"),
            PatternFill("solid", fgColor="FFF2CC"),
            Alignment(horizontal="center", vertical="center", wrapText=True),
            Border(left=thin, right=thin, top=thin, bottom=thin),
        ),
        "header": (
            Font(bold=True, size=11, color="1F4E78"),
            PatternFill("solid", fgColor="DDEBF7"),
            Alignment(horizontal="center", vertical="center"),
            Border(left=thin, right=thin, top=thin, bottom=thin),
        ),
        "notice": (
            Font(italic=True, size=10, color="FF0000"),
            PatternFill("solid", fgColor="FCE4D6"),
            Alignment(horizontal="right", vertical="bottom"),
            Border(left=thin, right=thin, top=thin, bottom=thin),
        ),
        "body": (
            Font(size=11),
            PatternFill(fill_type=None),
            Alignment(horizontal="left", vertical="center"),
            Border(left=thin, right=thin, top=thin, bottom=thin),
        ),
    }


def _apply_style(ws, coord: str, style) -> None:
    font, fill, alignment, border = style
    cell = ws[coord]
    cell.font = font
    cell.fill = fill
    cell.alignment = alignment
    cell.border = border


def _set_widths(ws, widths: dict[int, float]) -> None:
    from openpyxl.utils import get_column_letter

    for index, width in widths.items():
        ws.column_dimensions[get_column_letter(index)].width = width


def _inject_cached_values(path: Path, sheet_values: dict[str, dict[str, Any]]) -> None:
    """向 XLSX worksheet XML 写入公式缓存值（模拟 Excel 保存的结果）。"""
    with zipfile.ZipFile(path, "r") as archive:
        files = {name: archive.read(name) for name in archive.namelist()}

    workbook_root = ET.fromstring(files["xl/workbook.xml"])
    rels_root = ET.fromstring(files["xl/_rels/workbook.xml.rels"])
    targets = {rel.get("Id"): rel.get("Target") for rel in rels_root.findall(f"{{{_PKG_REL_NS}}}Relationship")}
    xml_paths: dict[str, str] = {}
    for sheet in workbook_root.find(f"{{{_MAIN_NS}}}sheets"):
        name = sheet.get("name")
        rel_id = sheet.get(f"{{{_REL_NS}}}id")
        target = targets.get(rel_id, "")
        if target.startswith("/"):
            xml_paths[name] = target.lstrip("/")
        elif target.startswith("xl/"):
            xml_paths[name] = target
        else:
            xml_paths[name] = "xl/" + target

    for sheet_name, values in sheet_values.items():
        xml_path = xml_paths[sheet_name]
        root = ET.fromstring(files[xml_path])
        cells = {cell.get("r"): cell for cell in root.iter(f"{{{_MAIN_NS}}}c")}
        for coord, value in values.items():
            cell = cells.get(coord)
            if cell is None:
                raise RuntimeError(f"XLSX 缓存值注入失败：{sheet_name}!{coord} 不存在")
            if cell.find(f"{{{_MAIN_NS}}}f") is None:
                raise RuntimeError(f"XLSX 缓存值注入失败：{sheet_name}!{coord} 不是公式单元格")
            value_node = cell.find(f"{{{_MAIN_NS}}}v")
            if value_node is None:
                value_node = ET.SubElement(cell, f"{{{_MAIN_NS}}}v")
            if isinstance(value, bool):
                cell.set("t", "b")
                value_node.text = "1" if value else "0"
            elif isinstance(value, (int, float)):
                if "t" in cell.attrib:
                    del cell.attrib["t"]
                value_node.text = repr(value)
            else:
                cell.set("t", "str")
                value_node.text = str(value)
        files[xml_path] = ET.tostring(root, encoding="utf-8", xml_declaration=True)

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in files.items():
            archive.writestr(name, payload)


def _base_sheet(ws, *, title: str = "设备定值单") -> None:
    styles = _styles()
    ws["A1"] = title
    ws.merge_cells("A1:D1")
    _apply_style(ws, "A1", styles["title"])
    ws.row_dimensions[1].height = 32
    _set_widths(ws, {1: 24, 2: 18, 3: 14, 4: 10})


# ---------------------------------------------------------------------------
# XLSX 多 Sheet 样本
# ---------------------------------------------------------------------------


def _fill_basic_sheet(ws, *, version: str, rated_voltage) -> None:
    styles = _styles()
    _base_sheet(ws)
    rows = [
        ("A2", "版本号", "B2", version),
        ("A3", "设备名称", "B3", "主变保护装置"),
        ("A4", "额定电压(kV)", "B4", rated_voltage),
        ("A5", "定值", "B5", 1.5),
        ("A6", "投运日期", "B6", _dt.datetime(2024, 5, 6)),
        ("A7", "备注", "B7", "合并单元格内容"),
    ]
    for label_coord, label, value_coord, value in rows:
        ws[label_coord] = label
        _apply_style(ws, label_coord, styles["header"])
        ws[value_coord] = value
        _apply_style(ws, value_coord, styles["body"])
    ws["B2"].font = ws["B2"].font.copy(italic=True)
    ws["B2"].alignment = ws["B2"].alignment.copy(horizontal="left")
    ws["B4"].number_format = "0.00"
    ws["B5"].number_format = "0.00"
    ws["B5"].alignment = ws["B5"].alignment.copy(horizontal="right")
    ws["B6"].number_format = "yyyy-mm-dd"
    ws["B7"].font = ws["B7"].font.copy(bold=True)
    ws.merge_cells("B7:C7")
    _apply_style(ws, "C7", styles["body"])
    _apply_style(ws, "C2", styles["notice"])
    ws["C2"] = "示例样式"
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[7].height = 26


def _fill_params_sheet(ws) -> None:
    styles = _styles()
    _base_sheet(ws, title="参数表")
    entries = [
        ("A2", "定值项", styles["header"]),
        ("B2", "整定值", styles["header"]),
        ("A3", "过流I段", styles["body"]),
        ("B3", 1.2, styles["body"]),
        ("A4", "过流II段", styles["body"]),
        ("B4", 0.8, styles["body"]),
        ("A5", "零序", styles["body"]),
        ("B5", 0.05, styles["body"]),
    ]
    for coord, value, style in entries:
        ws[coord] = value
        _apply_style(ws, coord, style)
    ws["B3"].number_format = "0.00"
    ws["B4"].number_format = "0.00"
    ws["B5"].number_format = "0.000"
    ws.merge_cells("A6:B6")
    ws["A6"] = "说明：本表两侧完全一致"
    _apply_style(ws, "A6", styles["notice"])
    _set_widths(ws, {1: 26, 2: 20})
    ws.row_dimensions[6].height = 28


def _fill_only_sheet(ws, *, side_label: str) -> None:
    styles = _styles()
    _base_sheet(ws, title=side_label)
    ws["A2"] = "仅此侧存在的工作表"
    ws["B2"] = 1
    _apply_style(ws, "A2", styles["body"])
    _apply_style(ws, "B2", styles["body"])


def generate_xlsx_fixtures(dest: Path = FIXTURES_DIR) -> list[Path]:
    from openpyxl import Workbook

    dest.mkdir(parents=True, exist_ok=True)
    left_path = dest / "sample_left.xlsx"
    right_path = dest / "sample_right.xlsx"

    left = Workbook()
    left.remove(left.active)
    left_basic = left.create_sheet("基本信息")
    _fill_basic_sheet(left_basic, version="V1.0.0", rated_voltage=220)
    left_params = left.create_sheet("参数表")
    _fill_params_sheet(left_params)
    left_only = left.create_sheet("仅左侧")
    _fill_only_sheet(left_only, side_label="仅左侧工作表的标题")
    left.save(left_path)

    right = Workbook()
    right.remove(right.active)
    right_params = right.create_sheet("参数表")
    _fill_params_sheet(right_params)
    right_basic = right.create_sheet("基本信息")
    _fill_basic_sheet(right_basic, version="V1.0.1", rated_voltage="220")
    right_only = right.create_sheet("仅右侧")
    _fill_only_sheet(right_only, side_label="仅右侧工作表的标题")
    right.save(right_path)
    return [left_path, right_path]


def _fill_formula_sheet(ws, *, formula_for_b2: str) -> None:
    styles = _styles()
    _base_sheet(ws, title="公式与缓存值")
    ws["A1"] = 2
    ws["A2"] = 3
    ws["B1"] = "=A1+A2"
    ws["B2"] = formula_for_b2
    ws["B3"] = "=A1-A2"
    ws["B4"] = '=IF(A1>1,"是","否")'
    ws["B5"] = "=A1/A2"
    ws["B5"].number_format = "0.00"
    for row in range(1, 6):
        for col in ("A", "B"):
            cell = ws[f"{col}{row}"]
            cell.font = styles["body"][0]
    ws["A1"].font = styles["header"][0]
    ws["B1"].font = styles["header"][0]
    _set_widths(ws, {1: 20, 2: 26})
    ws.row_dimensions[1].height = 22


def generate_formula_fixtures(dest: Path = FIXTURES_DIR) -> list[Path]:
    from openpyxl import Workbook

    dest.mkdir(parents=True, exist_ok=True)
    left_path = dest / "formula_left.xlsx"
    right_path = dest / "formula_right.xlsx"

    left = Workbook()
    _fill_formula_sheet(left.active, formula_for_b2="=A1*A2")
    left.save(left_path)
    # 左侧 B3 故意不注入缓存值，模拟“未在 Excel 中计算过”的公式。
    _inject_cached_values(
        left_path,
        {"公式与缓存值": {"B1": 5, "B2": 6, "B4": "是", "B5": 0.6666666666666666}},
    )

    right = Workbook()
    _fill_formula_sheet(right.active, formula_for_b2="=A1*A2*1")
    right.save(right_path)
    _inject_cached_values(
        right_path,
        {"公式与缓存值": {"B1": 5, "B2": 6, "B3": -1, "B4": "是", "B5": 0.67}},
    )
    return [left_path, right_path]


# ---------------------------------------------------------------------------
# XLS 样本
# ---------------------------------------------------------------------------


def generate_xls_fixtures(dest: Path = FIXTURES_DIR) -> list[Path]:
    import xlwt

    dest.mkdir(parents=True, exist_ok=True)
    left_path = dest / "sample_left.xls"
    right_path = dest / "sample_right.xls"
    styles = {
        "header": xlwt.easyxf(
            "font: bold on, colour_index blue; pattern: pattern solid, fore_colour ice_blue; "
            "align: horiz center, vert center;"
        ),
        "body": xlwt.easyxf("align: vert center;"),
        "notice": xlwt.easyxf(
            "font: italic on, colour_index red; pattern: pattern solid, fore_colour light_orange; "
            "align: horiz right;"
        ),
        "date": xlwt.easyxf("align: horiz right;", num_format_str="YYYY-MM-DD"),
    }

    def build(path: Path, *, version: str, third_value: float) -> None:
        workbook = xlwt.Workbook(encoding="utf-8")
        sheet = workbook.add_sheet("XLS示例")
        headers = ["设备", "版本", "定值", "日期", "合并"]
        for col, header in enumerate(headers):
            sheet.write(0, col, header, styles["header"])
        sheet.write(1, 0, "主变", styles["body"])
        sheet.write(1, 1, version, styles["notice"])
        sheet.write(1, 2, 12.5, styles["body"])
        sheet.write(1, 3, _dt.datetime(2024, 5, 6), styles["date"])
        sheet.write_merge(1, 1, 4, 5, "合并值", styles["body"])
        sheet.write(2, 0, "线路", styles["body"])
        sheet.write(2, 1, "V2.0", styles["body"])
        sheet.write(2, 2, third_value, styles["body"])
        sheet.write(2, 3, _dt.datetime(2024, 6, 1), styles["date"])
        for col, width in {0: 20, 1: 12, 2: 10, 3: 16, 4: 18, 5: 18}.items():
            sheet.col(col).width = int(width * 256)
        sheet.row(0).height = 30 * 20
        sheet.row(1).height = 24 * 20
        second = workbook.add_sheet("相同表")
        second.write(0, 0, "键", styles["header"])
        second.write(0, 1, "值", styles["header"])
        second.write(1, 0, "A", styles["body"])
        second.write(1, 1, 1.0, styles["body"])
        workbook.save(str(path))

    build(left_path, version="V1.0", third_value=0.85)
    build(right_path, version="V1.1", third_value=0.86)
    return [left_path, right_path]


# ---------------------------------------------------------------------------
# CSV 样本
# ---------------------------------------------------------------------------


_CSV_LEFT = (
    '设备名称,版本,说明,数值\n'
    '"主变,一号",V1.0.0,"第一行\n第二行",1.50\n'
    '线路保护,V2.0.0,正常,2\n'
    '\n'
    '含引号字段,"他说""你好""",末行,3.00\n'
)

_CSV_RIGHT = (
    '设备名称,版本,说明,数值\n'
    '"主变,一号",V1.0.1,"第一行\n第二行已修改",1.50\n'
    '线路保护,V2.0.0,正常,2\n'
    '\n'
    '含引号字段,"他说""你好""",末行,3.10\n'
)

_CSV_SEMICOLON = (
    "名称;版本;备注\n"
    "主变;V1.0;\"含分号;的字段\"\n"
    "线路;V2.0;正常\n"
)

_CSV_GBK = "设备,版本,备注\n主变,V9.9,中文编码测试\n"

_CSV_RAGGED = "A,B,C\n1,2\n3,4,5,6\n"


def generate_csv_fixtures(dest: Path = FIXTURES_DIR) -> list[Path]:
    dest.mkdir(parents=True, exist_ok=True)
    outputs: list[Path] = []
    for name, content in (
        ("sample_left.csv", _CSV_LEFT),
        ("sample_right.csv", _CSV_RIGHT),
        ("sample_semicolon.csv", _CSV_SEMICOLON),
        ("sample_ragged.csv", _CSV_RAGGED),
    ):
        path = dest / name
        path.write_text(content, encoding="utf-8")
        outputs.append(path)
    gbk_path = dest / "sample_gbk.csv"
    gbk_path.write_bytes(_CSV_GBK.encode("gb18030"))
    outputs.append(gbk_path)
    return outputs


# ---------------------------------------------------------------------------
# 统一入口
# ---------------------------------------------------------------------------


def generate_all(dest: Path = FIXTURES_DIR) -> list[Path]:
    outputs: list[Path] = []
    outputs.extend(generate_xlsx_fixtures(dest))
    outputs.extend(generate_formula_fixtures(dest))
    outputs.extend(generate_xls_fixtures(dest))
    outputs.extend(generate_csv_fixtures(dest))
    return outputs


_FIXTURE_NAMES = [
    "sample_left.xlsx",
    "sample_right.xlsx",
    "formula_left.xlsx",
    "formula_right.xlsx",
    "sample_left.xls",
    "sample_right.xls",
    "sample_left.csv",
    "sample_right.csv",
    "sample_semicolon.csv",
    "sample_gbk.csv",
    "sample_ragged.csv",
]


def ensure_fixtures(dest: Path = FIXTURES_DIR) -> Path:
    """样本缺失时自动生成（测试入口使用）。"""
    dest.mkdir(parents=True, exist_ok=True)
    if not all((dest / name).exists() for name in _FIXTURE_NAMES):
        generate_all(dest)
    return dest


if __name__ == "__main__":
    for generated in generate_all(FIXTURES_DIR):
        print(generated)
