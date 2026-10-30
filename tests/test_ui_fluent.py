"""Fluent 设计令牌、深浅色与导航/批量矩阵的 UI 冒烟测试。"""

from __future__ import annotations

from pathlib import Path

import pytest

from engineerdiff.models import AppSettings
from engineerdiff.ui import theme


def test_theme_tokens_and_stylesheet():
    assert theme.LIGHT.dark is False
    assert theme.DARK.dark is True
    # 差异红/待复核橙在深浅主题都必须有独立语义色
    for p in (theme.LIGHT, theme.DARK):
        assert p.diff_red and p.diff_red_bg
        assert p.warning and p.warning_bg
        assert p.diff_red != p.warning
        qss = theme.stylesheet(p)
        assert p.accent in qss
        assert "QTabBar#modeNav" in qss


def test_theme_apply_switches_palette(qapp):
    old = theme.palette()
    try:
        theme.apply_theme("dark", qapp)
        assert theme.palette().name == "dark"
        theme.apply_theme("light", qapp)
        assert theme.palette().name == "light"
    finally:
        theme.apply_theme(old.name, qapp)


@pytest.mark.gui
def test_main_navigation_and_dark_theme(qapp):
    from engineerdiff.ui.main_window import MainWindow

    theme.apply_theme("light", qapp)
    win = MainWindow()
    win.show()
    qapp.processEvents()
    assert win.pages.count() == 2
    assert win.nav.count() == 2
    win.nav.setCurrentIndex(1)
    qapp.processEvents()
    assert win.pages.currentIndex() == 1
    theme.apply_theme("dark", qapp)
    qapp.processEvents()
    assert theme.palette().name == "dark"
    assert win.palette().color(win.palette().ColorRole.Window).lightness() < 100
    win.close()


@pytest.mark.gui
def test_batch_excel_matrix_groups_by_sheet(qapp, tmp_path: Path):
    import openpyxl

    from engineerdiff.batch import run_batch
    from engineerdiff.ui.batch_dialog import BatchCompareDialog

    def make(path: Path, version: str, sheet2: bool) -> Path:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "参数"
        ws["A1"] = "版本"
        ws["B1"] = version
        wb.create_sheet("第二页")["A1"] = "基准" if sheet2 else "候选"
        wb.save(path)
        return path

    base = make(tmp_path / "a.xlsx", "V1", True)
    cand = make(tmp_path / "b.xlsx", "V2", False)
    result = run_batch(str(base), [str(cand)], AppSettings())
    dlg = BatchCompareDialog(None, AppSettings())
    dlg._populate(result)
    assert dlg.table.rowCount() >= len(result.rows) + 2
    assert len(dlg._groups) == 2
    # 组行跨越全部列
    assert dlg.table.columnSpan(dlg._groups[0][0], 0) == dlg.table.columnCount()
    dlg.close()


def test_settings_theme_default_is_system():
    assert AppSettings().theme == "system"
