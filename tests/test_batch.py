"""批量比较引擎与对话框测试。"""

from __future__ import annotations

import time
from pathlib import Path

import pytest

from engineerdiff.batch import BatchError, run_batch
from engineerdiff.models import AppSettings, Severity


def _write(path: Path, text: str) -> Path:
    path.write_text(text, encoding="utf-8")
    return path


def test_batch_text_or_semantics(tmp_path: Path):
    base = _write(tmp_path / "base.txt", "版本 V1.0\n定值 1.50\n备注 x\n")
    b = _write(tmp_path / "b.txt", "版本 V1.1\n定值 1.50\n备注 x\n")
    c = _write(tmp_path / "c.txt", "版本 V1.0\n定值 1.80\n备注 y\n")
    result = run_batch(str(base), [str(b), str(c)], AppSettings())
    assert result.mode.value == "text"
    assert result.file_count == 3
    assert result.different_count == 3
    rows = {row.location: row for row in result.rows}
    assert rows["第 1 行"].values == ["版本 V1.1", "版本 V1.0"]
    assert rows["第 1 行"].differs == [True, False]
    assert rows["第 2 行"].values == ["定值 1.50", "定值 1.80"]
    assert rows["第 2 行"].differs == [False, True]
    assert rows["第 3 行"].differs == [False, True]
    assert all(row.severity == Severity.CONFIRMED for row in result.rows)


def test_batch_text_candidate_only_line(tmp_path: Path):
    base = _write(tmp_path / "base.txt", "a\nb\n")
    c = _write(tmp_path / "c.txt", "a\nnew\nb\n")
    result = run_batch(str(base), [str(c)])
    assert result.different_count == 1
    row = result.rows[0]
    assert row.baseline == "—"
    assert row.values == ["new"]
    assert row.differs == [True]


def test_batch_excel_or_cells(tmp_path: Path):
    import openpyxl

    def make(path: Path, version: str, value: float, extra_sheet: bool = False) -> Path:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "基本信息"
        ws["A1"] = "版本"
        ws["B1"] = version
        ws["A2"] = "定值"
        ws["B2"] = value
        if extra_sheet:
            wb.create_sheet("仅基准")["A1"] = "x"
        wb.save(path)
        return path

    base = make(tmp_path / "base.xlsx", "V1.0", 1.5, True)
    b = make(tmp_path / "b.xlsx", "V1.1", 1.5)
    c = make(tmp_path / "c.xlsx", "V1.0", 1.8)
    result = run_batch(str(base), [str(b), str(c)])
    assert result.mode.value == "excel"
    rows = {row.location: row for row in result.rows}
    assert "基本信息!B1" in rows
    assert rows["基本信息!B1"].values == ["V1.1", "V1.0"]
    assert rows["基本信息!B1"].differs == [True, False]
    assert rows["基本信息!B2"].differs == [False, True]
    assert any("工作表" in loc for loc in rows)
    assert result.different_count == 3


def test_batch_rejects_mixed_types(tmp_path: Path):
    base = _write(tmp_path / "base.txt", "a\n")
    image = tmp_path / "a.png"
    image.write_bytes(b"not really png")
    # 扩展名先决定类型；混合类型必须在真正读取前被拒绝
    with pytest.raises(BatchError):
        run_batch(str(base), [str(image)])


@pytest.mark.ocr
def test_batch_image_ocr_rows(tmp_path: Path):
    from engineerdiff.ocr.engine import OcrService
    from tests.assets.generate_images import render_ui_screenshot

    a = render_ui_screenshot(tmp_path / "a.png")["path"]
    b = render_ui_screenshot(tmp_path / "b.png", values={"version": "1.2.4"})["path"]
    c = render_ui_screenshot(tmp_path / "c.png", values={"value_in_1": "1.80"})["path"]
    svc = OcrService(tmp_path / "ocr-cache")
    result = run_batch(str(a), [str(b), str(c)], AppSettings(), ocr_service=svc)
    assert result.mode.value == "image"
    assert result.different_count >= 2
    # 版本差异：至少 b 一列与基准不同；数值差异：至少 c 一列不同
    for row in result.rows:
        if "1.2.3" in row.baseline or "1.2.4" in " ".join(row.values):
            assert row.differs[0] is True
        if "1.50" in row.baseline or "1.80" in " ".join(row.values):
            assert row.differs[1] is True


@pytest.mark.gui
def test_batch_dialog_populates_table(qapp, tmp_path: Path):
    from engineerdiff.ui.batch_dialog import BatchCompareDialog

    base = _write(tmp_path / "base.txt", "版本 V1.0\n定值 1.50\n")
    b = _write(tmp_path / "b.txt", "版本 V1.1\n定值 1.50\n")
    c = _write(tmp_path / "c.txt", "版本 V1.0\n定值 1.80\n")
    dlg = BatchCompareDialog(None, AppSettings(), str(base), str(b))
    dlg._add_paths([str(c)])
    dlg.show()
    qapp.processEvents()
    assert dlg.file_list.count() == 3
    dlg._start()
    deadline = time.time() + 20
    while dlg._worker is not None and dlg._worker.isRunning() and time.time() < deadline:
        qapp.processEvents()
        time.sleep(0.02)
    qapp.processEvents()
    assert dlg.table.rowCount() == 2
    assert dlg.table.columnCount() == 4
    assert dlg.table.horizontalHeaderItem(1).text().startswith("基准：")
    dlg.close()
